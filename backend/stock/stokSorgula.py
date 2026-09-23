import re
import time
from multiprocessing import Pool, cpu_count

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


store_map = {
    "331": "İnternet Mağazası",
    "252": "IKEA Ümraniye",
    "194": "IKEA Bayrampaşa",
    "530": "IKEA Kartal",
}
store_codes = list(store_map.keys())

green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
orange_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

stock_api_url = "https://www.ikea.com.tr/_ws/general.aspx/CheckStock"
product_url_template = "https://www.ikea.com.tr/p{}"
request_headers = {
    "Content-Type": "application/json",
    "User-Agent": "Mozilla/5.0",
}


def normalize_stock_code(stock_code):
    return re.sub(r"\D", "", str(stock_code).strip())


def format_stock_code(stock_code):
    query_code = normalize_stock_code(stock_code)
    if len(query_code) == 8:
        return f"{query_code[:3]}.{query_code[3:6]}.{query_code[6:]}"
    return str(stock_code).strip()


def parse_title(html_text):
    match = re.search(r"<title>(.*?)</title>", html_text, flags=re.I | re.S)
    if not match:
        return "Ürün adı bulunamadı"

    title = re.sub(r"\s+", " ", match.group(1)).strip()
    return title.replace(" - IKEA", "")


def parse_price(html_text):
    patterns = [
        r'"price"\s*:\s*"?([0-9]+(?:[.,][0-9]+)?)',
        r'"Price"\s*:\s*"?([0-9]+(?:[.,][0-9]+)?)',
        r'"salesPrice"\s*:\s*"?([0-9]+(?:[.,][0-9]+)?)',
    ]

    for pattern in patterns:
        match = re.search(pattern, html_text, flags=re.I)
        if match:
            return float(match.group(1).replace(",", "."))

    return None


def parse_image_url(html_text):
    patterns = [
        r'"image"\s*:\s*"([^"]+)"',
        r'<meta property="og:image" content="([^"]+)"',
        r'<meta name="twitter:image" content="([^"]+)"',
    ]

    for pattern in patterns:
        match = re.search(pattern, html_text, flags=re.I)
        if match:
            return match.group(1).replace("\\/", "/")

    return None


def get_product_info(query_code):
    r = requests.get(
        product_url_template.format(query_code),
        headers=request_headers,
        timeout=15,
    )
    r.raise_for_status()
    return {
        "product_name": parse_title(r.text),
        "unit_price": parse_price(r.text),
        "image_url": parse_image_url(r.text),
    }


def get_product_name(query_code):
    return get_product_info(query_code)["product_name"]


def get_product_price(query_code):
    return get_product_info(query_code)["unit_price"]


def get_store_status(query_code, store_code):
    payload = {"stockCode": query_code, "storeCode": store_code}
    r = requests.post(
        stock_api_url,
        json=payload,
        headers=request_headers,
        timeout=15,
    )
    r.raise_for_status()

    data = r.json().get("d", {}).get("Data", {})
    status = data.get("Status", "")

    if status == "InStock":
        return "Var"
    if status == "CriticalStock":
        return "Kritik"
    if status == "InvalidStock":
        return "Hata"
    return "Yok"


def check_stock(stock_code):
    return check_stock_data(stock_code, include_product_info=True)


def check_stock_data(stock_code, include_product_info=True):
    original_code = str(stock_code).strip()
    query_code = normalize_stock_code(original_code)

    if not query_code:
        return {
            "product_name": "HATA",
            "stock_code": original_code,
            "query_code": query_code,
            "unit_price": None,
            "image_url": None,
            "statuses": {store_map[code]: "Hata" for code in store_codes},
        }

    product_name = format_stock_code(original_code)
    unit_price = None
    image_url = None
    if include_product_info:
        try:
            product_info = get_product_info(query_code)
            product_name = product_info["product_name"]
            unit_price = product_info["unit_price"]
            image_url = product_info["image_url"]
        except Exception:
            product_name = "Ürün adı alınamadı"

    statuses = {}
    for store_code in store_codes:
        try:
            status = get_store_status(query_code, store_code)
        except Exception:
            status = "Hata"
        statuses[store_map[store_code]] = status
        time.sleep(0.2)

    return {
        "product_name": product_name,
        "stock_code": format_stock_code(original_code),
        "query_code": query_code,
        "unit_price": unit_price,
        "image_url": image_url,
        "statuses": statuses,
    }


def process_stock(stock_code):
    original_code = str(stock_code).strip()
    print(f"\n{original_code} kontrol başlıyor")

    result = check_stock(original_code)
    print(f"Ürün adı: {result['product_name']}")
    for store_name, status in result["statuses"].items():
        print(f"   {store_name}: {status}")

    if result["product_name"] == "HATA":
        return ["HATA", original_code] + ["Hata"] * len(store_codes)

    return [
        result["product_name"],
        result["stock_code"],
        *[result["statuses"][store_map[code]] for code in store_codes],
    ]


def main():
    start_time = time.time()

    input_file = "stok_listesi.xlsx"
    sheet_name = "Sayfa1"
    output_file = "stock_kontrol.xlsx"

    wb_in = load_workbook(input_file)
    ws_in = wb_in[sheet_name]
    stock_codes = [row[0].value for row in ws_in.iter_rows(min_row=2) if row[0].value]

    worker_count = min(4, cpu_count(), len(stock_codes) or 1)
    print(f"Başladı. Toplam ürün: {len(stock_codes)} | Paralel işlem: {worker_count}")

    with Pool(worker_count) as pool:
        results = pool.map(process_stock, stock_codes)

    wb = Workbook()
    ws = wb.active
    ws.title = "Stok Durumu"

    headers = ["Ürün Adı", "StockCode"] + [store_map[code] for code in store_codes]
    ws.append(headers)

    for row in results:
        ws.append(row)
        current_row = ws.max_row

        if row[0] == "HATA":
            for col in range(1, len(row) + 1):
                ws.cell(row=current_row, column=col).fill = gray_fill
            continue

        for i, status in enumerate(row[2:], start=3):
            cell = ws.cell(row=current_row, column=i)
            if status == "Var":
                cell.fill = green_fill
            elif status == "Kritik":
                cell.fill = orange_fill
            elif status == "Hata":
                cell.fill = gray_fill
            else:
                cell.fill = red_fill

    wb.save(output_file)
    elapsed = time.time() - start_time
    avg = elapsed / len(stock_codes) if stock_codes else 0
    print(f"\nBitti. Dosya: {output_file}")
    print(f"Toplam süre: {elapsed:.2f} saniye | Ortalama: {avg:.2f} s/ürün")


if __name__ == "__main__":
    main()
