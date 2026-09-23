import json
import os
import re
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from io import BytesIO
from collections import OrderedDict
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import requests
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from stokSorgula import check_stock_data


BASE_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = BASE_DIR / "templates"
OUTPUTS_DIR = BASE_DIR / "outputs"
OZON_TEMPLATE_PATH = TEMPLATES_DIR / "Tüm depolardaki stokları güncelleme şablonu_19.09.2026.xlsx"
DEPOT_TYPE_PATH = TEMPLATES_DIR / "Depo Tipi.xlsx"
SETTINGS_PATH = BASE_DIR / "settings.json"
PRODUCT_CODE_RE = re.compile(r"(?<!\d)(\d{3}\.?\d{3}\.?\d{2})(?!\d)")
QUANTITY_RE = re.compile(r"(?:\t|\s+)(\d+)\s*$")
STORE_ORDER = ["IKEA Ümraniye", "IKEA Kartal", "IKEA Bayrampaşa", "İnternet Mağazası"]
FAST_STORES = ["IKEA Ümraniye", "IKEA Kartal"]
SLOW_STORES = ["IKEA Bayrampaşa", "İnternet Mağazası"]
DEFAULT_SETTINGS = {
    "in_stock_quantity": 10,
    "critical_stock_quantity": 2,
    "fast_stores": FAST_STORES,
    "slow_stores": SLOW_STORES,
}
OZON_DEPOTS = {
    ("UMT ExtraSmall", "5"): "UMT ExtraSmall - 5 Gün (1020005006595190)",
    ("UMT Express TR", "5"): "UMT Express TR - 5 Gün (1020005006595230)",
    ("UMT ExtraSmall", "2"): "UMT ExtraSmall - 2 Gün (1020005007889530)",
    ("UMT Express TR", "2"): "UMT Express TR - 2 Gün (1020005007890390)",
}
STOCK_TEMPLATE_JOBS = {}
STOCK_TEMPLATE_JOBS_LOCK = threading.Lock()
JOB_RETENTION_SECONDS = 60 * 60


def register_pdf_fonts():
    regular = Path("C:/Windows/Fonts/arial.ttf")
    bold = Path("C:/Windows/Fonts/arialbd.ttf")
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont("ArialTR", str(regular)))
        pdfmetrics.registerFont(TTFont("ArialTR-Bold", str(bold)))
        return "ArialTR", "ArialTR-Bold"
    return "Helvetica", "Helvetica-Bold"


PDF_FONT, PDF_BOLD_FONT = register_pdf_fonts()


def normalize_settings(raw_settings=None):
    raw_settings = raw_settings or {}
    settings = dict(DEFAULT_SETTINGS)

    for key in ("in_stock_quantity", "critical_stock_quantity"):
        try:
            value = int(raw_settings.get(key, settings[key]))
        except (TypeError, ValueError):
            raise ValueError(f"{key} sayısal olmalı")
        if value < 0:
            raise ValueError(f"{key} negatif olamaz")
        settings[key] = value

    for key in ("fast_stores", "slow_stores"):
        stores = raw_settings.get(key, settings[key])
        if not isinstance(stores, list):
            raise ValueError(f"{key} liste olmalı")
        clean_stores = []
        for store in stores:
            if store not in STORE_ORDER:
                raise ValueError(f"Geçersiz mağaza: {store}")
            if store not in clean_stores:
                clean_stores.append(store)
        if not clean_stores:
            raise ValueError(f"{key} boş olamaz")
        settings[key] = clean_stores

    return settings


def load_settings():
    if not SETTINGS_PATH.exists():
        return normalize_settings(DEFAULT_SETTINGS)

    with SETTINGS_PATH.open("r", encoding="utf-8") as file:
        return normalize_settings(json.load(file))


def save_settings(payload):
    settings = normalize_settings(payload)
    SETTINGS_PATH.write_text(
        json.dumps(settings, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return settings


def parse_order_text(text):
    items = OrderedDict()
    skipped_lines = []

    for line_no, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue

        code_match = PRODUCT_CODE_RE.search(line)
        if not code_match:
            skipped_lines.append({"line": line_no, "text": raw_line, "reason": "Ürün kodu bulunamadı"})
            continue

        raw_code = code_match.group(1)
        query_code = re.sub(r"\D", "", raw_code)
        display_code = f"{query_code[:3]}.{query_code[3:6]}.{query_code[6:]}"

        without_code = (line[: code_match.start()] + " " + line[code_match.end() :]).strip()
        quantity_match = QUANTITY_RE.search(without_code)
        quantity = int(quantity_match.group(1)) if quantity_match else 1
        name = QUANTITY_RE.sub("", without_code).strip(" -\t") or display_code

        if query_code not in items:
            items[query_code] = {
                "stock_code": display_code,
                "entered_name": name,
                "quantity": 0,
                "lines": [],
            }

        items[query_code]["quantity"] += quantity
        items[query_code]["lines"].append(line_no)
        if items[query_code]["entered_name"] == display_code and name != display_code:
            items[query_code]["entered_name"] = name

    return list(items.values()), skipped_lines


def build_order_item(item, include_product_info):
    stock = check_stock_data(item["stock_code"], include_product_info=include_product_info)
    unit_price = stock.get("unit_price")
    line_total = unit_price * item["quantity"] if unit_price is not None else None
    available_stores = [
        store_name
        for store_name, status in stock["statuses"].items()
        if status in {"Var", "Kritik"}
    ]
    return {
        **item,
        "product_name": stock["product_name"],
        "unit_price": unit_price,
        "line_total": line_total,
        "image_url": stock.get("image_url"),
        "statuses": stock["statuses"],
        "available_stores": available_stores,
    }


def check_order(text, include_product_info=True):
    items, skipped_lines = parse_order_text(text)
    results = [None] * len(items)
    worker_count = min(16, max(1, len(items)))

    with ThreadPoolExecutor(max_workers=worker_count) as executor:
        future_map = {
            executor.submit(build_order_item, item, include_product_info): index
            for index, item in enumerate(items)
        }
        for future in as_completed(future_map):
            results[future_map[future]] = future.result()

    return {"items": results, "skipped_lines": skipped_lines}


def normalize_code(stock_code):
    return re.sub(r"\D", "", str(stock_code or "").strip())


def display_code(stock_code):
    query_code = normalize_code(stock_code)
    if len(query_code) == 8:
        return f"{query_code[:3]}.{query_code[3:6]}.{query_code[6:]}"
    return str(stock_code or "").strip()


def extract_codes(text):
    return [normalize_code(match) for match in PRODUCT_CODE_RE.findall(str(text or ""))]


def validate_depot_type_workbook(path=DEPOT_TYPE_PATH):
    if not path.exists():
        raise FileNotFoundError(f"{path.name} bulunamadı")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    issues = []
    exact_rows = {}
    code_types = {}
    valid_types = {"UMT ExtraSmall", "UMT Express TR"}
    row_count = 0

    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw_product, depot_type = (row + (None, None))[:2] if isinstance(row, tuple) else (None, None)
        if not raw_product and not depot_type:
            continue

        row_count += 1
        product_text = str(raw_product or "").strip()
        clean_type = str(depot_type or "").strip()
        codes = extract_codes(product_text)

        if not codes:
            issues.append({"level": "error", "line": row_index, "text": product_text, "reason": "Ürün kodu bulunamadı"})
            continue
        if clean_type not in valid_types:
            issues.append({"level": "error", "line": row_index, "text": product_text, "reason": "Depo tipi geçersiz"})
            continue

        exact_key = tuple(codes)
        if exact_key in exact_rows:
            issues.append({
                "level": "warning",
                "line": row_index,
                "text": product_text,
                "reason": f"Aynı ürün kombinasyonu daha önce {exact_rows[exact_key]}. satırda var",
            })
        else:
            exact_rows[exact_key] = row_index

        for code in codes:
            previous_type = code_types.get(code)
            if previous_type and previous_type != clean_type:
                issues.append({
                    "level": "warning",
                    "line": row_index,
                    "text": product_text,
                    "reason": f"{display_code(code)} kodu farklı depo tiplerinde kullanılıyor",
                })
            code_types[code] = clean_type

    return {
        "row_count": row_count,
        "unique_products": len(exact_rows),
        "unique_codes": len(code_types),
        "issue_count": len(issues),
        "error_count": sum(1 for issue in issues if issue["level"] == "error"),
        "warning_count": sum(1 for issue in issues if issue["level"] == "warning"),
        "issues": issues,
    }


def parse_stock_template_text(text):
    products = OrderedDict()
    skipped_lines = []

    for line_no, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue

        codes = extract_codes(line)
        if not codes:
            skipped_lines.append({"line": line_no, "text": raw_line, "reason": "Ürün kodu bulunamadı"})
            continue

        key = tuple(codes)
        if key not in products:
            products[key] = {
                "codes": codes,
                "article": line,
                "source_text": line,
                "lines": [],
            }
        products[key]["lines"].append(line_no)

    return list(products.values()), skipped_lines


def load_depot_products(path=DEPOT_TYPE_PATH):
    if not path.exists():
        raise FileNotFoundError(f"{path.name} bulunamadı")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    products = OrderedDict()
    skipped_lines = []

    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw_product, depot_type = (row + (None, None))[:2] if isinstance(row, tuple) else (None, None)
        if not raw_product and not depot_type:
            continue

        codes = extract_codes(raw_product)
        clean_type = str(depot_type or "").strip()
        if not codes:
            skipped_lines.append({"line": row_index, "text": str(raw_product or ""), "reason": "Ürün kodu bulunamadı"})
            continue
        if clean_type not in {"UMT ExtraSmall", "UMT Express TR"}:
            skipped_lines.append({"line": row_index, "text": str(raw_product or ""), "reason": "Depo tipi geçersiz"})
            continue

        key = tuple(codes)
        if key not in products:
            product_text = str(raw_product).strip()
            products[key] = {
                "codes": codes,
                "article": product_text,
                "source_text": product_text,
                "depot_type": clean_type,
                "lines": [],
            }
        products[key]["lines"].append(row_index)

    return list(products.values()), skipped_lines


def load_depot_types(path=DEPOT_TYPE_PATH):
    if not path.exists():
        raise FileNotFoundError(f"{path.name} bulunamadı")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    by_exact_codes = {}
    by_single_code = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_product, depot_type = (row + (None, None))[:2] if isinstance(row, tuple) else (None, None)
        if not raw_product or not depot_type:
            continue

        codes = extract_codes(raw_product)
        clean_type = str(depot_type).strip()
        if not codes or clean_type not in {"UMT ExtraSmall", "UMT Express TR"}:
            continue

        by_exact_codes[tuple(codes)] = clean_type
        for code in codes:
            by_single_code[code] = clean_type

    return by_exact_codes, by_single_code


def resolve_depot_type(product, by_exact_codes, by_single_code):
    exact = by_exact_codes.get(tuple(product["codes"]))
    if exact:
        return exact

    component_types = {by_single_code.get(code) for code in product["codes"]}
    component_types.discard(None)
    if len(component_types) == 1:
        return component_types.pop()
    return None


def component_group_status(result, stores):
    statuses = result.get("statuses") or {}
    group_statuses = [statuses.get(store, "Hata") for store in stores]
    if "Var" in group_statuses:
        return "Var"
    if "Kritik" in group_statuses:
        return "Kritik"
    return "Yok"


def group_stock_quantity(component_results, stores, settings):
    per_component = []
    for result in component_results:
        per_component.append(component_group_status(result, stores))

    if per_component and all(status == "Var" for status in per_component):
        return settings["in_stock_quantity"]
    if per_component and all(status in {"Var", "Kritik"} for status in per_component):
        return settings["critical_stock_quantity"]
    return 0


def build_component_report(product, component_results, selected_depot, selected_qty, settings):
    components = []
    for code, result in zip(product["codes"], component_results):
        statuses = result.get("statuses") or {}
        components.append({
            "code": display_code(code),
            "fast_status": component_group_status(result, settings["fast_stores"]),
            "slow_status": component_group_status(result, settings["slow_stores"]),
            "stores": {store: statuses.get(store, "Hata") for store in STORE_ORDER},
        })

    fast_blockers = [item["code"] for item in components if item["fast_status"] == "Yok"]
    slow_blockers = [item["code"] for item in components if item["slow_status"] == "Yok"]
    if selected_depot:
        reason = f"{selected_depot} için {selected_qty} adet yazıldı"
    else:
        reason = "Uygun depo grubu bulunamadı; tüm depolar sıfırlandı"

    return {
        "article": product["article"],
        "depot_type": product.get("depot_type") or "",
        "selected_depot": selected_depot,
        "selected_qty": selected_qty,
        "component_count": len(product["codes"]),
        "components": components,
        "fast_blockers": fast_blockers,
        "slow_blockers": slow_blockers,
        "reason": reason,
    }


def update_job(job_id, **values):
    with STOCK_TEMPLATE_JOBS_LOCK:
        job = STOCK_TEMPLATE_JOBS.get(job_id)
        if job:
            job.update(values)
            job["updated_at"] = time.time()


def get_job(job_id):
    with STOCK_TEMPLATE_JOBS_LOCK:
        return STOCK_TEMPLATE_JOBS.get(job_id)


def cleanup_jobs():
    cutoff = time.time() - JOB_RETENTION_SECONDS
    with STOCK_TEMPLATE_JOBS_LOCK:
        old_ids = [
            job_id
            for job_id, job in STOCK_TEMPLATE_JOBS.items()
            if job.get("status") in {"done", "error"} and job.get("updated_at", 0) < cutoff
        ]
        for job_id in old_ids:
            STOCK_TEMPLATE_JOBS.pop(job_id, None)


def public_job_status(job):
    payload = {
        "id": job["id"],
        "status": job["status"],
        "stage": job.get("stage", ""),
        "message": job.get("message", ""),
        "total": job.get("total", 0),
        "completed": job.get("completed", 0),
        "percent": job.get("percent", 0),
        "generated_rows": job.get("generated_rows", 0),
        "skipped_count": len(job.get("skipped_lines", [])),
        "error_count": len(job.get("errors", [])),
        "summary": job.get("summary", {}),
        "skipped_lines": job.get("skipped_lines", [])[:50],
        "errors": job.get("errors", [])[:50],
        "depot_validation": job.get("depot_validation", {}),
        "component_reports": job.get("component_reports", [])[:100],
        "history_file": job.get("history_file", ""),
        "settings": job.get("settings", load_settings()),
    }
    if job.get("status") == "error":
        payload["error"] = job.get("error", "Bilinmeyen hata")
    return payload


def build_stock_template_rows_from_products(products, skipped_lines=None, job_id=None, settings=None):
    settings = settings or load_settings()
    skipped_lines = list(skipped_lines or [])
    by_exact_codes, by_single_code = load_depot_types()
    all_codes = sorted({code for product in products for code in product["codes"]})
    stock_results = {}
    errors = []
    if job_id:
        update_job(
            job_id,
            stage="stock",
            message=f"IKEA stokları sorgulanıyor: 0 / {len(all_codes)}",
            total=len(all_codes),
            completed=0,
            percent=0,
        )

    with ThreadPoolExecutor(max_workers=min(8, max(1, len(all_codes)))) as executor:
        future_map = {
            executor.submit(check_stock_data, display_code(code), False): code
            for code in all_codes
        }
        completed = 0
        for future in as_completed(future_map):
            code = future_map[future]
            try:
                stock_results[code] = future.result()
            except Exception as exc:
                errors.append({"code": display_code(code), "error": str(exc)})
                stock_results[code] = {
                    "stock_code": display_code(code),
                    "query_code": code,
                    "statuses": {store: "Hata" for store in STORE_ORDER},
                }
            completed += 1
            if job_id:
                percent = int((completed / max(1, len(all_codes))) * 85)
                update_job(
                    job_id,
                    completed=completed,
                    percent=percent,
                    message=f"IKEA stokları sorgulanıyor: {completed} / {len(all_codes)}",
                )

    rows = []
    prepared = []
    component_reports = []
    for product in products:
        depot_type = product.get("depot_type") or resolve_depot_type(product, by_exact_codes, by_single_code)
        if not depot_type:
            skipped_lines.append({
                "line": ", ".join(str(line) for line in product["lines"]),
                "text": product["source_text"],
                "reason": "Depo tipi bulunamadı",
            })
            continue

        component_results = [stock_results[code] for code in product["codes"]]
        fast_qty = group_stock_quantity(component_results, settings["fast_stores"], settings)
        slow_qty = group_stock_quantity(component_results, settings["slow_stores"], settings)
        selected_depot = None
        selected_qty = 0
        if fast_qty > 0:
            selected_depot = OZON_DEPOTS[(depot_type, "2")]
            selected_qty = fast_qty
        elif slow_qty > 0:
            selected_depot = OZON_DEPOTS[(depot_type, "5")]
            selected_qty = slow_qty

        product_rows = [
            [
                depot_name,
                product["article"],
                "",
                selected_qty if depot_name == selected_depot else 0,
            ]
            for depot_name in OZON_DEPOTS.values()
        ]
        rows.extend(product_rows)
        if len(product["codes"]) > 1:
            component_reports.append(build_component_report(product, component_results, selected_depot, selected_qty, settings))
        prepared.append({
            "article": product["article"],
            "depot_type": depot_type,
            "fast_qty": fast_qty,
            "slow_qty": slow_qty,
            "selected_depot": selected_depot,
            "selected_qty": selected_qty,
            "component_count": len(product["codes"]),
        })

    positive_products = sum(1 for item in prepared if item["selected_qty"] > 0)
    zero_products = sum(1 for item in prepared if item["selected_qty"] == 0)

    return {
        "rows": rows,
        "prepared": prepared,
        "summary": {
            "product_count": len(prepared),
            "generated_rows": len(rows),
            "positive_products": positive_products,
            "zero_products": zero_products,
            "zeroed_rows": len(rows) - positive_products,
        },
        "settings": settings,
        "component_reports": component_reports,
        "skipped_lines": skipped_lines,
        "errors": errors,
    }


def build_stock_template_rows(text):
    products, skipped_lines = parse_stock_template_text(text)
    return build_stock_template_rows_from_products(products, skipped_lines)


def save_output_history(workbook_data):
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_path = OUTPUTS_DIR / f"ozon-stok-sablonu-{timestamp}.xlsx"
    output_path.write_bytes(workbook_data)
    return output_path


def run_stock_template_job(job_id):
    try:
        cleanup_jobs()
        settings = load_settings()
        update_job(job_id, status="running", stage="load", message="Depo tipi listesi okunuyor.", percent=0)
        depot_validation = validate_depot_type_workbook()
        products, skipped_lines = load_depot_products()
        update_job(
            job_id,
            message=f"{len(products)} ürün hazırlandı. IKEA sorgusu başlıyor.",
            total=len({code for product in products for code in product["codes"]}),
            completed=0,
            depot_validation=depot_validation,
        )
        template_data = build_stock_template_rows_from_products(products, skipped_lines, job_id=job_id, settings=settings)
        update_job(job_id, stage="workbook", message="Ozon Excel şablonu hazırlanıyor.", percent=92)
        workbook_data = build_ozon_stock_workbook(template_data["rows"])
        history_path = save_output_history(workbook_data)
        update_job(
            job_id,
            status="done",
            stage="done",
            message="Şablon hazır.",
            percent=100,
            generated_rows=len(template_data["rows"]),
            summary=template_data["summary"],
            settings=settings,
            depot_validation=depot_validation,
            component_reports=template_data["component_reports"],
            skipped_lines=template_data["skipped_lines"],
            errors=template_data["errors"],
            history_file=str(history_path),
            file=workbook_data,
        )
    except Exception as exc:
        update_job(job_id, status="error", stage="error", message=str(exc), error=str(exc), percent=100)


def start_stock_template_job():
    cleanup_jobs()
    job_id = uuid.uuid4().hex
    now = time.time()
    with STOCK_TEMPLATE_JOBS_LOCK:
        STOCK_TEMPLATE_JOBS[job_id] = {
            "id": job_id,
            "status": "queued",
            "stage": "queued",
            "message": "İş sıraya alındı.",
            "total": 0,
            "completed": 0,
            "percent": 0,
            "generated_rows": 0,
            "summary": {},
            "settings": load_settings(),
            "depot_validation": {},
            "component_reports": [],
            "history_file": "",
            "skipped_lines": [],
            "errors": [],
            "created_at": now,
            "updated_at": now,
            "file": None,
        }
    threading.Thread(target=run_stock_template_job, args=(job_id,), daemon=True).start()
    return STOCK_TEMPLATE_JOBS[job_id]


def cell_text(ref, value, style_id):
    cell = ET.Element("c", {"r": ref, "s": str(style_id), "t": "inlineStr"})
    inline = ET.SubElement(cell, "is")
    text = ET.SubElement(inline, "t")
    text.text = str(value)
    return cell


def cell_number(ref, value, style_id):
    cell = ET.Element("c", {"r": ref, "s": str(style_id)})
    number = ET.SubElement(cell, "v")
    number.text = str(int(value))
    return cell


def build_ozon_stock_workbook(rows):
    if not OZON_TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"{OZON_TEMPLATE_PATH.name} bulunamadı")

    ET.register_namespace("", "http://schemas.openxmlformats.org/spreadsheetml/2006/main")
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    ET.register_namespace("mc", "http://schemas.openxmlformats.org/markup-compatibility/2006")
    ET.register_namespace("x14ac", "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac")
    ET.register_namespace("x14", "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main")
    ET.register_namespace("xm", "http://schemas.microsoft.com/office/excel/2006/main")
    ET.register_namespace("xr", "http://schemas.microsoft.com/office/spreadsheetml/2014/revision")
    ET.register_namespace("xr2", "http://schemas.microsoft.com/office/spreadsheetml/2015/revision2")
    ET.register_namespace("xr3", "http://schemas.microsoft.com/office/spreadsheetml/2016/revision3")
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    row_count = max(1000, len(rows) + 1)

    with ZipFile(OZON_TEMPLATE_PATH, "r") as source:
        sheet_xml = source.read("xl/worksheets/sheet2.xml")
        root = ET.fromstring(sheet_xml)
        dimension = root.find("m:dimension", ns)
        if dimension is not None:
            dimension.set("ref", f"A1:E{row_count}")

        sheet_data = root.find("m:sheetData", ns)
        header_row = sheet_data.find("m:row[@r='1']", ns)
        sheet_data.clear()
        if header_row is not None:
            sheet_data.append(header_row)

        for row_index in range(2, row_count + 1):
            row_el = ET.Element("row", {
                "r": str(row_index),
                "spans": "1:5",
                "ht": "15.75",
                "customHeight": "1",
            })
            data_index = row_index - 2
            if data_index < len(rows):
                depot_name, article, product_name, quantity = rows[data_index]
                row_el.append(cell_text(f"A{row_index}", depot_name, 20))
                row_el.append(cell_text(f"B{row_index}", article, 20))
                row_el.append(cell_text(f"C{row_index}", product_name, 20))
                row_el.append(cell_number(f"D{row_index}", quantity, 25))
            else:
                for column, style_id in [("A", 20), ("B", 20), ("C", 20), ("D", 25), ("E", 22)]:
                    row_el.append(ET.Element("c", {"r": f"{column}{row_index}", "s": str(style_id)}))
            sheet_data.append(row_el)

        updated_sheet_text = ET.tostring(root, encoding="unicode", xml_declaration=True)
        if "xmlns:xr2=" not in updated_sheet_text:
            updated_sheet_text = updated_sheet_text.replace(
                "<worksheet ",
                '<worksheet xmlns:xr2="http://schemas.microsoft.com/office/spreadsheetml/2015/revision2" ',
                1,
            )
        if "xmlns:xr3=" not in updated_sheet_text:
            updated_sheet_text = updated_sheet_text.replace(
                "<worksheet ",
                '<worksheet xmlns:xr3="http://schemas.microsoft.com/office/spreadsheetml/2016/revision3" ',
                1,
            )
        updated_sheet = updated_sheet_text.encode("utf-8")
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = updated_sheet if item.filename == "xl/worksheets/sheet2.xml" else source.read(item.filename)
                target.writestr(item, data)

    return output.getvalue()


def money(value):
    if value is None:
        return "-"
    try:
        value = float(value)
    except (TypeError, ValueError):
        return "-"

    text = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if text.endswith(",00"):
        text = text[:-3]
    return f"{text} TL"


def status_color(status):
    if status == "Var":
        return colors.HexColor("#d8f0df")
    if status == "Kritik":
        return colors.HexColor("#fff0bf")
    if status == "Yok":
        return colors.HexColor("#ffd9d9")
    return colors.HexColor("#e5e7eb")


def pdf_product_image(image_url, cache):
    if not image_url:
        return ""

    try:
        if image_url not in cache:
            response = requests.get(image_url, timeout=10, headers={"User-Agent": "Mozilla/5.0"})
            response.raise_for_status()
            cache[image_url] = response.content

        image = Image(BytesIO(cache[image_url]), width=1.25 * cm, height=1.25 * cm)
        image.hAlign = "CENTER"
        return image
    except Exception:
        return ""


def build_order_pdf(items, stores):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=0.7 * cm,
        rightMargin=0.7 * cm,
        topMargin=0.7 * cm,
        bottomMargin=0.7 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleTR",
        parent=styles["Title"],
        fontName=PDF_BOLD_FONT,
        fontSize=15,
        leading=18,
        spaceAfter=6,
    )
    small_style = ParagraphStyle(
        "SmallTR",
        parent=styles["Normal"],
        fontName=PDF_FONT,
        fontSize=7,
        leading=8,
    )
    header_style = ParagraphStyle(
        "HeaderTR",
        parent=small_style,
        fontName=PDF_BOLD_FONT,
        alignment=TA_CENTER,
    )
    right_style = ParagraphStyle(
        "RightTR",
        parent=small_style,
        alignment=TA_RIGHT,
    )
    center_style = ParagraphStyle(
        "CenterTR",
        parent=small_style,
        alignment=TA_CENTER,
    )

    stores = stores or STORE_ORDER
    total_amount = sum(float(item.get("line_total") or 0) for item in items)
    total_qty = sum(int(item.get("quantity") or 0) for item in items)

    story = [
        Paragraph("Sipariş Stok Kontrol", title_style),
        Paragraph(
            f"Ürün çeşidi: {len(items)} &nbsp;&nbsp; Toplam adet: {total_qty} &nbsp;&nbsp; Toplam tutar: {money(total_amount)}",
            small_style,
        ),
        Spacer(1, 0.25 * cm),
    ]

    image_cache = {}
    headers = ["Görsel", "Ürün adı", "Kod", "Adet", "Birim", "Toplam", *stores]
    table_data = [[Paragraph(header, header_style) for header in headers]]

    for item in items:
        row = [
            pdf_product_image(item.get("image_url"), image_cache),
            Paragraph(str(item.get("product_name") or ""), small_style),
            Paragraph(str(item.get("stock_code") or ""), center_style),
            Paragraph(str(item.get("quantity") or ""), center_style),
            Paragraph(money(item.get("unit_price")), right_style),
            Paragraph(money(item.get("line_total")), right_style),
        ]
        statuses = item.get("statuses") or {}
        row.extend(Paragraph(str(statuses.get(store, "Hata")), center_style) for store in stores)
        table_data.append(row)

    page_width = landscape(A4)[0] - doc.leftMargin - doc.rightMargin
    fixed_width = 1.55 * cm + 2.6 * cm + 1.1 * cm + 1.7 * cm + 1.9 * cm + (len(stores) * 2.15 * cm)
    product_width = max(6.8 * cm, page_width - fixed_width)
    col_widths = [1.55 * cm, product_width, 2.6 * cm, 1.1 * cm, 1.7 * cm, 1.9 * cm] + [2.15 * cm] * len(stores)

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style = TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), PDF_FONT),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1ed")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1f2933")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cfd6ce")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
    )

    for row_index, item in enumerate(items, start=1):
        statuses = item.get("statuses") or {}
        if row_index % 2 == 0:
            style.add("BACKGROUND", (0, row_index), (5, row_index), colors.HexColor("#fbfbfa"))
        for store_index, store in enumerate(stores, start=6):
            style.add("BACKGROUND", (store_index, row_index), (store_index, row_index), status_color(statuses.get(store)))

    table.setStyle(style)
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


class SiparisHandler(BaseHTTPRequestHandler):
    def do_OPTIONS(self):
        self.send_response(204)
        self.send_cors_headers()
        self.end_headers()

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/api/health":
            self.send_json({"ok": True})
            return
        if path in {"/", "/index.html"}:
            self.send_file(BASE_DIR / "siparis.html", "text/html; charset=utf-8")
            return
        if path == "/api/settings":
            self.send_json({"settings": load_settings(), "stores": STORE_ORDER})
            return
        if path.startswith("/api/stock-template-status/"):
            job_id = path.rsplit("/", 1)[-1]
            job = get_job(job_id)
            if not job:
                self.send_json({"error": "İş bulunamadı"}, status=404)
                return
            self.send_json(public_job_status(job))
            return
        if path.startswith("/api/stock-template-download/"):
            job_id = path.rsplit("/", 1)[-1]
            job = get_job(job_id)
            if not job:
                self.send_json({"error": "İş bulunamadı"}, status=404)
                return
            if job.get("status") != "done" or not job.get("file"):
                self.send_json({"error": "Dosya henüz hazır değil"}, status=409)
                return
            self.send_bytes(
                job["file"],
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                'attachment; filename="ozon-stok-sablonu.xlsx"',
            )
            return
        self.send_json({"error": "Sayfa bulunamadı"}, status=404)

    def do_POST(self):
        path = urlparse(self.path).path
        if path not in {
            "/api/check-order",
            "/api/export-pdf",
            "/api/export-stock-template",
            "/api/start-stock-template",
            "/api/settings",
        }:
            self.send_json({"error": "Endpoint bulunamadı"}, status=404)
            return

        try:
            content_length = int(self.headers.get("Content-Length", "0"))
            body = self.rfile.read(content_length).decode("utf-8")
            payload = json.loads(body or "{}")
            if path == "/api/settings":
                self.send_json({"settings": save_settings(payload.get("settings", payload)), "stores": STORE_ORDER})
                return

            if path == "/api/start-stock-template":
                job = start_stock_template_job()
                self.send_json(public_job_status(job))
                return

            if path == "/api/check-order":
                text = payload.get("text", "")
                include_product_info = bool(payload.get("includeProductInfo", True))
                self.send_json(check_order(text, include_product_info=include_product_info))
                return

            if path == "/api/export-stock-template":
                text = payload.get("text", "")
                template_data = build_stock_template_rows(text)
                workbook_data = build_ozon_stock_workbook(template_data["rows"])
                self.send_bytes(
                    workbook_data,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    'attachment; filename="ozon-stok-sablonu.xlsx"',
                    extra_headers={
                        "X-Generated-Rows": str(len(template_data["rows"])),
                        "X-Skipped-Count": str(len(template_data["skipped_lines"])),
                        "X-Stock-Error-Count": str(len(template_data["errors"])),
                    },
                )
                return

            items = payload.get("items", [])
            stores = payload.get("stores", STORE_ORDER)
            pdf_data = build_order_pdf(items, stores)
            self.send_bytes(
                pdf_data,
                "application/pdf",
                'attachment; filename="siparis-stok.pdf"',
            )
        except Exception as exc:
            self.send_json({"error": str(exc)}, status=500)

    def send_file(self, path, content_type):
        if not path.exists():
            self.send_json({"error": f"{path.name} bulunamadı"}, status=404)
            return

        data = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.send_cors_headers()
        self.end_headers()
        self.wfile.write(data)

    def send_json(self, payload, status=200):
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_cors_headers()
        self.end_headers()
        self.wfile.write(data)

    def send_bytes(self, data, content_type, content_disposition=None, status=200, extra_headers=None):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        if content_disposition:
            self.send_header("Content-Disposition", content_disposition)
        if extra_headers:
            for key, value in extra_headers.items():
                self.send_header(key, value)
        self.send_cors_headers()
        self.end_headers()
        self.wfile.write(data)

    def send_cors_headers(self):
        origin = self.headers.get("Origin")
        self.send_header("Access-Control-Allow-Origin", origin or "*")
        self.send_header("Access-Control-Allow-Credentials", "true")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def log_message(self, format, *args):
        return


def main():
    port = int(os.environ.get("OMEGA_STOCK_PORT", "8010"))
    server = ThreadingHTTPServer(("127.0.0.1", port), SiparisHandler)
    print(f"OMEGA stok servisi: http://127.0.0.1:{port}")
    server.serve_forever()


if __name__ == "__main__":
    main()
